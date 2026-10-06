import asyncio, pandas as pd
import json
from io import BytesIO
from os import getenv
from zoneinfo import ZoneInfo

from aiogram.exceptions import TelegramNetworkError
from dateutil.relativedelta import relativedelta
from dotenv import load_dotenv
from datetime import datetime, timezone, timedelta
from aiogram import Bot, Dispatcher, F, types, Router
from aiogram.filters import CommandStart, Command
from aiogram.fsm.context import FSMContext
from aiogram.filters.state import State, StatesGroup
from aiogram.types import FSInputFile, InlineKeyboardButton, InlineKeyboardMarkup, CallbackQuery, BufferedInputFile
from openpyxl.workbook import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from model_ask import get_model_analysis, load_data
from channel_parse import parse_channel_web
from logging_config import setup_logging, logging
import re

# Инициализация логгирования
setup_logging()
logger = logging.getLogger(__name__)

# Инициализация секретов
load_dotenv()
TELEGRAM_BOT_API_TOKEN = getenv("TELEGRAM_BOT_API_TOKEN")

# Инициализация бота
bot = Bot(token=TELEGRAM_BOT_API_TOKEN)
dp = Dispatcher()

# Инициализация путей использования бота
user_router = Router()


# FSM состояние бота
class NavigateStates(StatesGroup):
    option_user_choice = State()
    get_time_choice = State()
    get_interval = State()
    get_channel_name = State()


class AnalyzeStates(StatesGroup):
    message_analyze = State()
    channel_analyze = State()
    get_analysis_results = State()


# Функции

# Безопасная отправка документа

async def safe_send_document(bot, **kwargs):
    for attempt in range(3):
        try:
            return await bot.send_document(
                **kwargs,
                request_timeout=30
            )
        except TelegramNetworkError as tg:
            logger.error(f"Ошибка: {tg}")
            if attempt == 2:
                raise
            await asyncio.sleep(2)


# Проверка валидности даты

async def validate_date_pair(date_string: str) -> dict:
    base_pattern = r'^(\d{2})\.(\d{2})\.(\d{4}) (\d{2})\.(\d{2})\.(\d{4})$'
    match = re.match(base_pattern, date_string)

    if not match:
        return {
            'valid': False,
            'message': "Неверный формат. Нужно: день.месяц.год день.месяц.год",
        }

    d1, m1, y1, d2, m2, y2 = map(int, match.groups())
    current_date = datetime.now()

    if not (1 <= m1 <= 12):
        return {
            'valid': False,
            'message': f"Неверный месяц в первой дате: {m1:02d}. Месяц должен быть от 01 до 12."
        }

    try:
        date1 = datetime(y1, m1, d1)
    except ValueError:
        if d1 < 1 or d1 > 31:
            msg = f"Неверный день в первой дате: {d1:02d}. День должен быть от 01 до 31."
        elif m1 == 2 and d1 > 29:
            msg = f"В феврале не может быть {d1:02d} дня. Проверьте високосный год."
        elif m1 in [4, 6, 9, 11] and d1 > 30:
            msg = f"В {m1:02d}-м месяце только 30 дней."
        else:
            msg = f"Некорректная первая дата: {d1:02d}.{m1:02d}.{y1:04d}"

        return {
            'valid': False,
            'message': msg
        }

    if date1 > current_date:
        return {
            'valid': False,
            'message': f"Первая дата ({date1.strftime('%d.%m.%Y')}) ещё не наступила!"
        }

    if not (1 <= m2 <= 12):
        return {
            'valid': False,
            'message': f"Неверный месяц во второй дате: {m2:02d}. Месяц должен быть от 01 до 12."
        }

    try:
        date2 = datetime(y2, m2, d2)
    except ValueError:
        if d2 < 1 or d2 > 31:
            msg = f"Неверный день во второй дате: {d2:02d}. День должен быть от 01 до 31."
        elif m2 == 2 and d2 > 29:
            msg = f"В феврале не может быть {d2:02d} дня. Проверьте високосный год."
        elif m2 in [4, 6, 9, 11] and d2 > 30:
            msg = f"В {m2:02d}-м месяце только 30 дней."
        else:
            msg = f"Некорректная вторая дата: {d2:02d}.{m2:02d}.{y2:04d}"

        return {
            'valid': False,
            'message': msg
        }

    if date2 > current_date:
        return {
            'valid': False,
            'message': f"Вторая дата ({date2.strftime('%d.%m.%Y')}) ещё не наступила!"
        }

    if date2 < date1:
        return {
            'valid': False,
            'message': (
                f"Вторая дата ({date2.strftime('%d.%m.%Y')}) "
                f"раньше первой ({date1.strftime('%d.%m.%Y')})!"
            )
        }

    return {
        'valid': True
    }


async def get_channel_analysis(channel_name, start_time, end_time, user_id) -> dict:
    if start_time and end_time:
        start_time = datetime.strptime(start_time, "%d.%m.%Y").replace(tzinfo=timezone.utc)
        end_time = datetime.strptime(end_time + " 23:59:59", "%d.%m.%Y %H:%M:%S").replace(tzinfo=timezone.utc)

    parse_response = parse_channel_web(channel_name, start_time, end_time)

    if not parse_response:
        logger.error("Ошибка парсинга")
        return {"completion": False, "error": "parsing"}

    datum = []
    for post in parse_response:
        # Определяем текст поста
        if isinstance(post, dict) and "content" in post:
            post_text = post["content"]
        else:
            # Для обратной совместимости со старым форматом
            post_text = post
        model_response = await get_model_analysis(post_text)

        if not model_response["error"]:
            answer_list = json.loads(model_response["model_answer"])
            violations_flag = any(answer["fine"] > 0 for answer in answer_list)
        else:
            logger.error(f"Ошибка при обработке запроса:\n\n{post_text}")
            answer_list = []
            violations_flag = False

        descriptions = ""
        recommendations = ""
        if violations_flag:
            answer_list = sorted(answer_list, key=lambda answer: answer["fine"], reverse=True)

        for i, answer in enumerate(answer_list):
            descriptions += f"{str(i + 1)}) Описание:\n" + answer[
                "description"] + "\n\nВозможный штраф: " + str(answer["fine"]) + "\n\n"
            recommendations += f"{str(i + 1)}) Рекомендация:\n" + answer["recommendation"] + "\n\n"

        cell = {
            "Текст публикации": post_text,
            "Наличие нарушений": "Присутствуют 🚨" if violations_flag == True else "Отсутствуют ✅",
            "Описание": descriptions,
            "Рекомендации": recommendations
        }
        datum.append(cell)

    df = pd.DataFrame(datum)

    wb = Workbook()
    ws = wb.active

    # Ограничиваем имя листа до 31 символа
    sheet_title = f"Анализ {channel_name}"[:31]
    ws.title = sheet_title

    filename = f"Анализ канала {channel_name}.xlsx"

    # Настройки стилей
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    cell_alignment = Alignment(
        vertical='top',
        horizontal='left',
        wrap_text=True
    )

    # Записываем заголовки
    for col_idx, column_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=str(column_name))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', horizontal='center')
        cell.border = thin_border

    # Записываем данные
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = cell_alignment
            cell.border = thin_border

    # Автоподбор ширины колонок с учетом переноса текста
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        # Находим максимальную длину текста в колонке
        for cell in column:
            if cell.value:
                # Учитываем перенос строк
                lines = str(cell.value).split('\n')
                max_line_length = max(len(line) for line in lines)
                max_length = max(max_length, max_line_length)

        # Устанавливаем ширину (ограничиваем максимум 50 символов)
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Устанавливаем высоту строк для лучшего отображения
    for row in ws.iter_rows(min_row=2):
        max_lines = 1
        for cell in row:
            if cell.value:
                lines = str(cell.value).count('\n') + 1
                max_lines = max(max_lines, lines)
        ws.row_dimensions[row[0].row].height = max_lines * 15

    # Замораживаем заголовки
    ws.freeze_panes = 'A2'

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    excel_file = BufferedInputFile(
        file=excel_buffer.read(),
        filename=filename
    )

    try:
        # Отправляем файл
        await safe_send_document(
            bot,
            chat_id=user_id,
            document=excel_file,
            caption=f"📊 Анализ постов канала {channel_name} ({datetime.now().strftime('%d.%m.%Y %H:%M')})"
        )
        return {"completion": True}
    except Exception as e:
        logger.error(f"Ошибка отправки файла: {e}")
        return {"completion": False, "error": "document_send"}
    finally:
        logger.info(f"Обработано {len(datum)} постов")


# Клавиатуры бота
def option_user_choice():
    keyboard_list = [
        [InlineKeyboardButton(text='Проверить сообщение 💬', callback_data='message_check')],
        [InlineKeyboardButton(text='Проверить канал 🗂', callback_data='channel_check')]
    ]
    keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_list)
    return keyboard


def option_user_go_back():
    keyboard_list = [
        [InlineKeyboardButton(text='Вернуться назад ↩️', callback_data='back')]
    ]
    keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_list)
    return keyboard


def option_user_choose_time_interval():
    keyboard_list = [
        [InlineKeyboardButton(text='Все посты в канале 📂', callback_data='all_posts')],
        [InlineKeyboardButton(text='Указать период вручную 📅', callback_data='interval')],
        [InlineKeyboardButton(text='За последние 7 дней 🗓', callback_data='week')],
        [InlineKeyboardButton(text='За последний месяц 🗓', callback_data='month')],
        [InlineKeyboardButton(text='За последний год 🗓', callback_data='year')],
        [InlineKeyboardButton(text='Вернуться назад ↩️', callback_data='back')]
    ]
    keyboard = InlineKeyboardMarkup(inline_keyboard=keyboard_list)
    return keyboard


# Команда

@user_router.message(Command("secret_admin_statistics_request"))
async def handle_admin_request(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    now = datetime.now(ZoneInfo("Europe/Moscow"))
    month = now.strftime("%B").lower()
    file_name = f"logs/{month}_logs_{now.year}.json"
    data = await load_data(file_name)
    summ_input_tokens, summ_output_tokens, summ_mc_time = 0, 0, 0
    for cell in data:
        summ_input_tokens += cell["input_tokens"]
        summ_output_tokens += cell["output_tokens"]
        summ_mc_time += cell["Время ответа в мс"]
    await message.answer(
        f"📊 _Cтатистика за последний календарный месяц_" + "\n\n" +
        f"Количество запросов к модели: *{len(data)}*\n" + f"Среднее количество входящих токенов: *{summ_input_tokens // len(data)}*\n" +
        f"Среднее количество исходящих токенов: *{summ_output_tokens // len(data)}*\n" + f"Среднее время ответа (мс): *{summ_mc_time // len(data)}*\n",
        parse_mode="Markdown"
    )
    try:
        # Отправляем файл
        await safe_send_document(bot,
                                 chat_id=message.from_user.id,
                                 document=FSInputFile(file_name),
                                 caption=f"📊 Статистика постов ({datetime.now().strftime('%d.%m.%Y %H:%M')})"
                                 )
    except Exception as e:
        logger.error(f"Ошибка: {e}")
        await asyncio.sleep(0.5)
        await message.answer("Отправить документ не удалось, попробуйте позже")
    finally:
        return


# Начальный блок сообщений с выбором варианта использования бота
@user_router.message(CommandStart())
async def handle_start_message(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    username = message.from_user.username
    welcome_message = await message.answer(
        f"Привет, *{username}*!\n\nЯ являются телеграмм-ботом компании *Systeme Electric*, созданным для проверки сообщений на *наличие рекламы.*\n\nБуду рад помочь 🤝",
        reply_markup=option_user_choice(), parse_mode="Markdown"
    )
    logger.info("User %s started conversation", message.from_user.id)
    await state.update_data(last_message=welcome_message)
    await state.set_state(NavigateStates.option_user_choice)


@user_router.message(F.text, NavigateStates.option_user_choice)
async def attention_message_to_user(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    await message.answer(f"Пожалуйста, выбери одну из опций выше.", parse_mode="Markdown")
    await state.set_state(NavigateStates.option_user_choice)


@user_router.callback_query(F.data == 'message_check', NavigateStates.option_user_choice)
async def option_message_check(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    logger.info("User %s started message analysis", call.from_user.id)
    await last_message.edit_text(text="Пришли мне *текст сообщения* 📝",
                                 reply_markup=option_user_go_back(), parse_mode="Markdown")
    await state.set_state(AnalyzeStates.message_analyze)


@user_router.callback_query(F.data == 'back', AnalyzeStates.message_analyze)
async def back_to_previous(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    username = message.from_user.username
    await last_message.edit_text(
        text=f"Привет, *{username}*!\n\nЯ являются телеграмм-ботом компании *Systeme Electric*, созданным для проверки сообщений на *наличие рекламы.*\n\nБуду рад помочь 🤝",
        reply_markup=option_user_choice(), parse_mode="Markdown")
    await state.set_state(NavigateStates.option_user_choice)


@user_router.message((F.text | F.caption), AnalyzeStates.message_analyze)
async def message_analyze(message: types.Message, state: FSMContext):
    data = await state.get_data()
    last_message = data.get("last_message")
    text = message.text or message.caption
    await last_message.edit_text(text="Пришли мне *текст сообщения* 📝",
                                 reply_markup=None, parse_mode="Markdown")
    await asyncio.sleep(0.5)
    ai_answer = await message.answer("_Искусственный интеллект обрабатывает ваш запрос.._", parse_mode="Markdown")
    model_response = await get_model_analysis(text)
    logger.info(text)
    if model_response["error"]:
        await ai_answer.edit_text(text="Пожалуйста, попробуйте позже.")
        logger.error(f"Ошибка обработки сообщения: {model_response.get("model_answer")}")
        return

    answer_list = json.loads(model_response["model_answer"])
    violations_flag = any(answer["fine"] > 0 for answer in answer_list)

    if violations_flag:
        answer_list = sorted(answer_list, key=lambda answer: answer["fine"], reverse=True)

    descriptions = ""
    for i, answer in enumerate(answer_list):
        descriptions += f"*{str(i + 1)}) Описание:*\n" + answer["description"] + "\n\n*Возможный штраф:* " + str(
            answer["fine"]) + "\n\n*Рекомендация:*\n" + answer["recommendation"] + "\n\n"

    violation_status = "Присутствуют 🚨" if violations_flag else "Отсутствуют ✅"

    try:
        final_answer_to_user = f"*Результаты анализа 🔍*\n\n*Нарушения: *{violation_status}\n\n{descriptions}"
        await ai_answer.edit_text(text=final_answer_to_user, parse_mode="Markdown")

    except Exception as e:
        logger.error(f"Ошибка: {e}")
        descriptions = ""
        for i, answer in enumerate(answer_list):
            descriptions += f"{str(i + 1)}) Описание:\n" + answer["description"] + "\n\nВозможный штраф: " + str(
                answer["fine"]) + "\n\nРекомендация:\n" + answer["recommendation"] + "\n\n"

        violation_status = "Присутствуют 🚨" if violations_flag else "Отсутствуют ✅"

        final_answer_to_user = f"Результаты анализа 🔍\n\nНарушения: {violation_status}\n\n{descriptions}"
        await ai_answer.edit_text(text=final_answer_to_user, parse_mode=None)
    logger.info(final_answer_to_user)
    await asyncio.sleep(0.5)
    last_message = await message.answer("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.update_data(last_message=last_message)
    await state.set_state(NavigateStates.option_user_choice)


@user_router.callback_query(F.data == 'channel_check', NavigateStates.option_user_choice)
async def option_channel_check(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    logger.info("User %s started channel analysis", call.from_user.id)
    await last_message.edit_text(text="Пришли мне *тег канала* в формате *@systemeelectric_official* 🔗",
                                 reply_markup=option_user_go_back(), parse_mode="Markdown")
    await state.set_state(NavigateStates.get_channel_name)


@user_router.callback_query(F.data == 'back', NavigateStates.get_channel_name)
async def back_to_previous(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    await last_message.edit_text("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.set_state(NavigateStates.option_user_choice)


@user_router.message(F.text, NavigateStates.get_channel_name)
async def check_bot_rights(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    text = message.text
    logger.info(text)
    if "https://t.me/" in text:
        text = text.replace("https://t.me/", "@")
    if text[0] != "@":
        await asyncio.sleep(0.5)
        await message.answer("Пожалуйста, пришли тег канала в формате *@systemeelectric_official* 🔗",
                             parse_mode="Markdown")
        logger.warning("Invalid channel name: %s", text)
        await state.set_state(NavigateStates.get_channel_name)
        return
    await last_message.edit_text(text=f"Канал для проверки: {text}", reply_markup=None)
    await asyncio.sleep(0.5)
    last_message = await message.answer("Какой период постов нужно проверить?",
                                        reply_markup=option_user_choose_time_interval())
    await state.update_data(channel_name=text, last_message=last_message)
    await state.set_state(NavigateStates.get_time_choice)


@user_router.callback_query(F.data == 'back', NavigateStates.get_time_choice)
async def back_to_previous(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    await last_message.edit_text("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.set_state(NavigateStates.option_user_choice)


@user_router.callback_query(F.data == 'all_posts', NavigateStates.get_time_choice)
async def get_all_timestamp(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    channel_name = data.get("channel_name")
    await last_message.edit_text(
        text=f"Ты выбрал анализ *всех постов* в канале. Он может занять *продолжительное время* ⏱️.",
        parse_mode="Markdown")
    await asyncio.sleep(0.5)
    ai_answer = await call.message.answer(
        "_Искусственный интеллект обрабатывает ваш запрос..._",
        parse_mode="Markdown")
    task_result_description = await get_channel_analysis(channel_name[1:], None, None, call.from_user.id)
    await ai_answer.delete()
    await asyncio.sleep(0.5)
    if not task_result_description.get("completion") and task_result_description.get("error") == "parsing":
        await call.message.answer("Пожалуйста, проверь не ошибся ли ты в теге или измени временной промежуток постов.")
    elif task_result_description.get("error") == "document_send":
        await call.message.answer("Пожалуйста, попробуйте позже, ошибка при отправке документа")
    last_message = await call.message.answer("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.update_data(last_message=last_message)
    await state.set_state(NavigateStates.option_user_choice)


@user_router.callback_query(F.data.in_({"week", "month", "year"}), NavigateStates.get_time_choice)
async def get_exact_timestamp(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    channel_name = data.get("channel_name")
    end_time = datetime.now()
    dict_dates = {
        "week": {
            "date": (end_time - timedelta(weeks=1)).strftime("%d.%m.%Y"),
            "message": f"Ты выбрал анализ постов за *последнюю неделю*."},
        "month": {
            "date": (end_time - relativedelta(months=1)).strftime("%d.%m.%Y"),
            "message": f"Ты выбрал анализ постов за *последний месяц*.\n\nОн может занять *продолжительное время* ⏱️."},
        "year": {
            "date": (end_time - relativedelta(years=1)).strftime("%d.%m.%Y"),
            "message": f"Ты выбрал анализ постов за *последний год*.\n\nОн может занять *продолжительное время* ⏱️."}
    }
    await last_message.edit_text(text=dict_dates[call.data]["message"], parse_mode="Markdown")
    await asyncio.sleep(0.5)
    ai_answer = await call.message.answer(
        "_Искусственный интеллект обрабатывает ваш запрос..._",
        parse_mode="Markdown")
    task_result_description = await get_channel_analysis(channel_name[1:], dict_dates[call.data]["date"],
                                                         end_time.strftime("%d.%m.%Y"), call.from_user.id)
    await ai_answer.delete()
    await asyncio.sleep(0.5)
    if not task_result_description.get("completion") and task_result_description.get("error") == "parsing":
        await call.message.answer("Пожалуйста, проверь не ошибся ли ты в теге или измени временной промежуток постов.")
    elif task_result_description.get("error") == "document_send":
        await call.message.answer("Пожалуйста, попробуйте позже, ошибка при отправке документа")
    last_message = await call.message.answer("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.update_data(last_message=last_message)
    await state.set_state(NavigateStates.option_user_choice)


@user_router.callback_query(F.data == 'interval', NavigateStates.get_time_choice)
async def get_timestamps_interval(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    await last_message.edit_text(
        text=f"Ты выбрал анализ постов *определенного временного промежутка*.\n\nПожалуйста, пришли мне *дату начала* и *дату конца интервала* в формате день.месяц.год .\n\n*Например:* 08.12.2025 26.12.2025",
        reply_markup=option_user_go_back(), parse_mode="Markdown")
    await state.set_state(NavigateStates.get_interval)


@user_router.callback_query(F.data == 'back', NavigateStates.get_interval)
async def back_to_previous(call: CallbackQuery, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    last_message = data.get("last_message")
    await last_message.edit_text("Какой период постов нужно проверить?",
                                 reply_markup=option_user_choose_time_interval())
    await state.set_state(NavigateStates.get_time_choice)


@user_router.message(F.text, NavigateStates.get_interval)
async def get_interval(message: types.Message, state: FSMContext):
    await asyncio.sleep(0.5)
    data = await state.get_data()
    channel_name = data.get("channel_name")
    last_message = data.get("last_message")
    valid_dict = await validate_date_pair(message.text)
    if not valid_dict.get("valid"):
        await message.answer(valid_dict.get("message"))
        await state.set_state(NavigateStates.get_interval)
        return
    await asyncio.sleep(0.5)
    await last_message.edit_text(text=f"Твой выбранный промежуток: *{message.text}*", reply_markup=None,
                                 parse_mode="Markdown")
    dates = message.text.split()
    ai_answer = await message.answer(
        "_Искусственный интеллект обрабатывает ваш запрос, он может занять продолжительное время..._",
        parse_mode="Markdown")
    task_result_description = await get_channel_analysis(channel_name[1:], dates[0], dates[1], message.from_user.id)
    await ai_answer.delete()
    await asyncio.sleep(0.5)
    if not task_result_description.get("completion") and task_result_description.get("error") == "parsing":
        await message.answer("Пожалуйста, проверь не ошибся ли ты в теге или измени временной промежуток постов.")
    elif task_result_description.get("error") == "document_send":
        await message.answer("Пожалуйста, попробуйте позже, ошибка при отправке документа")
    last_message = await message.answer("Выбери дальнейшую опцию", reply_markup=option_user_choice())
    await state.update_data(last_message=last_message)
    await state.set_state(NavigateStates.option_user_choice)


async def main():
    dp.include_router(user_router)
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())