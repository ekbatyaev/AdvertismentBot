import asyncio, pandas as pd
from io import BytesIO
from aiogram.exceptions import TelegramNetworkError
from datetime import datetime, timezone
from aiogram.types import BufferedInputFile
from openpyxl.workbook import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import re
from app.settings import logger
from app.bot.config import bot
from app.ai.functions.ad_analysis import analyze_text
from app.parser.telegram_web_parser import parse_channel_web


_PARSER_SEMAPHORE = asyncio.Semaphore(4)
# Безопасная отправка документа

async def safe_send_document(**kwargs):
    for attempt in range(3):
        try:
            return await bot.send_document(
                **kwargs,
                request_timeout=30
            )
        except TelegramNetworkError as tg:
            logger.error(f"Ошибка: {tg}")
            if attempt == 2:
                raise
            await asyncio.sleep(2)


# Проверка валидности даты

async def validate_date_pair(date_string: str) -> dict:
    base_pattern = r'^(\d{2})\.(\d{2})\.(\d{4}) (\d{2})\.(\d{2})\.(\d{4})$'
    match = re.match(base_pattern, date_string)

    if not match:
        return {
            'valid': False,
            'message': "Неверный формат. Нужно: день.месяц.год день.месяц.год",
        }

    d1, m1, y1, d2, m2, y2 = map(int, match.groups())
    current_date = datetime.now()

    if not (1 <= m1 <= 12):
        return {
            'valid': False,
            'message': f"Неверный месяц в первой дате: {m1:02d}. Месяц должен быть от 01 до 12."
        }

    try:
        date1 = datetime(y1, m1, d1)
    except ValueError:
        if d1 < 1 or d1 > 31:
            msg = f"Неверный день в первой дате: {d1:02d}. День должен быть от 01 до 31."
        elif m1 == 2 and d1 > 29:
            msg = f"В феврале не может быть {d1:02d} дня. Проверьте високосный год."
        elif m1 in [4, 6, 9, 11] and d1 > 30:
            msg = f"В {m1:02d}-м месяце только 30 дней."
        else:
            msg = f"Некорректная первая дата: {d1:02d}.{m1:02d}.{y1:04d}"

        return {
            'valid': False,
            'message': msg
        }

    if date1 > current_date:
        return {
            'valid': False,
            'message': f"Первая дата ({date1.strftime('%d.%m.%Y')}) ещё не наступила!"
        }

    if not (1 <= m2 <= 12):
        return {
            'valid': False,
            'message': f"Неверный месяц во второй дате: {m2:02d}. Месяц должен быть от 01 до 12."
        }

    try:
        date2 = datetime(y2, m2, d2)
    except ValueError:
        if d2 < 1 or d2 > 31:
            msg = f"Неверный день во второй дате: {d2:02d}. День должен быть от 01 до 31."
        elif m2 == 2 and d2 > 29:
            msg = f"В феврале не может быть {d2:02d} дня. Проверьте високосный год."
        elif m2 in [4, 6, 9, 11] and d2 > 30:
            msg = f"В {m2:02d}-м месяце только 30 дней."
        else:
            msg = f"Некорректная вторая дата: {d2:02d}.{m2:02d}.{y2:04d}"

        return {
            'valid': False,
            'message': msg
        }

    if date2 > current_date:
        return {
            'valid': False,
            'message': f"Вторая дата ({date2.strftime('%d.%m.%Y')}) ещё не наступила!"
        }

    if date2 < date1:
        return {
            'valid': False,
            'message': (
                f"Вторая дата ({date2.strftime('%d.%m.%Y')}) "
                f"раньше первой ({date1.strftime('%d.%m.%Y')})!"
            )
        }

    return {
        'valid': True
    }


async def get_channel_analysis(channel_name, start_time, end_time, user_id) -> dict:
    if start_time and end_time:
        start_time = datetime.strptime(start_time, "%d.%m.%Y").replace(tzinfo=timezone.utc)
        end_time = datetime.strptime(end_time + " 23:59:59", "%d.%m.%Y %H:%M:%S").replace(tzinfo=timezone.utc)

    async with _PARSER_SEMAPHORE:
        parse_response = await asyncio.to_thread(parse_channel_web, channel_name, start_time, end_time)

    if not parse_response:
        logger.error("Ошибка парсинга")
        return {"completion": False, "error": "parsing"}

    datum = []
    for post in parse_response:
        # Определяем текст поста
        if isinstance(post, dict) and "content" in post:
            post_text = post["content"]
        else:
            # Для обратной совместимости со старым форматом
            post_text = post

        answer_list = await analyze_text(text = post_text, temperature = 0.4)

        if answer_list:
            violations_flag = any(answer["fine"] > 0 for answer in answer_list)
        else:
            logger.error(f"Ошибка при обработке запроса:\n\n{post_text}")
            answer_list = []
            violations_flag = False
        descriptions = ""
        recommendations = ""
        if violations_flag:
            answer_list = sorted(answer_list, key=lambda answer: answer["fine"], reverse=True)

        for i, answer in enumerate(answer_list):
            descriptions += f"{str(i + 1)}) Описание:\n" + answer[
                "description"] + "\n\nВозможный штраф: " + str(answer["fine"]) + "\n\n"
            recommendations += f"{str(i + 1)}) Рекомендация:\n" + answer["recommendation"] + "\n\n"

        cell = {
            "Текст публикации": post_text,
            "Наличие нарушений": "Присутствуют 🚨" if violations_flag == True else "Отсутствуют ✅",
            "Описание": descriptions,
            "Рекомендации": recommendations
        }
        datum.append(cell)

    df = pd.DataFrame(datum)

    wb = Workbook()
    ws = wb.active

    # Ограничиваем имя листа до 31 символа
    sheet_title = f"Анализ {channel_name}"[:31]
    ws.title = sheet_title

    filename = f"Анализ канала {channel_name}.xlsx"

    # Настройки стилей
    header_font = Font(name='Arial', size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    cell_alignment = Alignment(
        vertical='top',
        horizontal='left',
        wrap_text=True
    )

    # Записываем заголовки
    for col_idx, column_name in enumerate(df.columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=str(column_name))
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', horizontal='center')
        cell.border = thin_border

    # Записываем данные
    for row_idx, row in enumerate(df.itertuples(index=False), 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = cell_alignment
            cell.border = thin_border

    # Автоподбор ширины колонок с учетом переноса текста
    for column in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column[0].column)

        # Находим максимальную длину текста в колонке
        for cell in column:
            if cell.value:
                # Учитываем перенос строк
                lines = str(cell.value).split('\n')
                max_line_length = max(len(line) for line in lines)
                max_length = max(max_length, max_line_length)

        # Устанавливаем ширину (ограничиваем максимум 50 символов)
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Устанавливаем высоту строк для лучшего отображения
    for row in ws.iter_rows(min_row=2):
        max_lines = 1
        for cell in row:
            if cell.value:
                lines = str(cell.value).count('\n') + 1
                max_lines = max(max_lines, lines)
        ws.row_dimensions[row[0].row].height = max_lines * 15

    # Замораживаем заголовки
    ws.freeze_panes = 'A2'

    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    excel_file = BufferedInputFile(
        file=excel_buffer.read(),
        filename=filename
    )

    try:
        # Отправляем файл
        await safe_send_document(
            chat_id=user_id,
            document=excel_file,
            caption=f"📊 Анализ постов канала {channel_name} ({datetime.now().strftime('%d.%m.%Y %H:%M')})"
        )
        return {"completion": True}
    except Exception as e:
        logger.error(f"Ошибка отправки файла: {e}")
        return {"completion": False, "error": "document_send"}
    finally:
        logger.info(f"Обработано {len(datum)} постов")